"""Excel (.xlsx) reading and writing via the optional ``openpyxl`` dependency."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING, List, Optional, Sequence, Union

from ._table import PathLike, Table, TablemdError
from ._text import display_width, stringify

if TYPE_CHECKING:  # pragma: no cover
    pass

__all__ = ["write_xlsx", "read_xlsx", "safe_sheet_name"]

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_MAX_COL_WIDTH = 60


def _require_openpyxl():
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise TablemdError(
            "openpyxl is required for Excel support: pip install 'tablemd[xlsx]'"
        ) from exc
    return openpyxl


def _clean(value):
    """Drop control characters Excel refuses to store (openpyxl would raise)."""
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # type: ignore

        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def safe_sheet_name(name: str, taken: Sequence[str] = ()) -> str:
    """Return *name* adjusted to Excel's sheet-name rules and unique within *taken*."""
    cleaned = _BAD_SHEET_CHARS.sub("_", name).strip().strip("'") or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    n = 2
    lowered = {t.lower() for t in taken}
    while candidate.lower() in lowered:
        suffix = f" ({n})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    return candidate


def write_xlsx(
    tables: Sequence[Table],
    path: PathLike,
    *,
    sheet_names: Optional[Sequence[str]] = None,
    infer_types: bool = True,
    bold_header: bool = True,
    freeze_header: bool = True,
    autofit: bool = True,
) -> None:
    """Write one or more tables to an ``.xlsx`` workbook, one sheet per table."""
    if not tables:
        raise TablemdError("no tables to write")
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Alignment as XLAlignment  # type: ignore
    from openpyxl.styles import Font  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    wb = openpyxl.Workbook()
    default_ws = wb.active
    used: List[str] = []

    for idx, table in enumerate(tables):
        wanted = (
            sheet_names[idx]
            if sheet_names is not None and idx < len(sheet_names)
            else f"Table{idx + 1}"
        )
        title = safe_sheet_name(wanted, used)
        used.append(title)
        ws = default_ws if idx == 0 else wb.create_sheet()
        ws.title = title

        ws.append([_clean(h) for h in table.header])
        body = table.typed_rows() if infer_types else table.rows
        for row in body:
            ws.append([_clean(v) for v in row])
        # openpyxl turns any string starting with "=" into a live formula;
        # table cells are data, so force them back to plain text.
        for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for c in row_cells:
                if c.data_type == "f":
                    c.data_type = "s"

        if bold_header:
            for cell in ws[1]:
                cell.font = Font(bold=True)
        if freeze_header:
            ws.freeze_panes = "A2"

        for col_idx, align in enumerate(table.align, start=1):
            if not align:
                continue
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2):
                for c in cell:
                    c.alignment = XLAlignment(horizontal=align)

        if autofit:
            for col_idx in range(1, table.ncols + 1):
                width = display_width(table.header[col_idx - 1])
                for row in table.rows:
                    width = max(width, max((display_width(part) for part in row[col_idx - 1].split("\n")), default=0))
                letter = get_column_letter(col_idx)
                ws.column_dimensions[letter].width = min(_MAX_COL_WIDTH, max(8, width + 2))

        if any("\n" in c for row in table.rows for c in row):
            for row_cells in ws.iter_rows(min_row=2):
                for c in row_cells:
                    if isinstance(c.value, str) and "\n" in c.value:
                        c.alignment = XLAlignment(
                            wrap_text=True,
                            horizontal=c.alignment.horizontal if c.alignment else None,
                        )

    wb.save(str(path))


def read_xlsx(
    path: PathLike,
    *,
    sheet: Union[str, int, None] = None,
) -> List[Table]:
    """Read an ``.xlsx`` workbook into tables, one per sheet.

    *sheet* selects a single worksheet by name or 0-based index; ``None``
    returns every non-empty sheet. The first row is taken as the header.
    """
    openpyxl = _require_openpyxl()
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types here
        raise TablemdError(f"cannot open {path!s} as an Excel workbook: {exc}") from exc
    try:
        if sheet is None:
            sheets = list(wb.worksheets)
        else:
            if isinstance(sheet, str) and sheet not in wb.sheetnames and sheet.strip().isdigit():
                sheet = int(sheet)  # "--sheet 2" on the command line
            if isinstance(sheet, int):
                try:
                    sheets = [wb.worksheets[sheet]]
                except IndexError as exc:
                    raise TablemdError(
                        f"workbook has no sheet index {sheet} (it has {len(wb.worksheets)} sheet(s), 0-based)"
                    ) from exc
            else:
                if sheet not in wb.sheetnames:
                    raise TablemdError(f"workbook has no sheet named {sheet!r}; available: {wb.sheetnames}")
                sheets = [wb[sheet]]

        tables: List[Table] = []
        for ws in sheets:
            rows = [[stringify(v) for v in r] for r in ws.iter_rows(values_only=True)]
            # Drop fully-empty trailing rows/columns that openpyxl may report.
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            while ncols and all(len(r) < ncols or r[ncols - 1] == "" for r in rows):
                ncols -= 1
            rows = [r[:ncols] + [""] * (ncols - len(r[:ncols])) for r in rows]
            if ncols == 0:
                continue
            tables.append(Table(header=rows[0], rows=rows[1:]))
        return tables
    finally:
        wb.close()
